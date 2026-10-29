import csv
from pathlib import Path

SUPPORTED = {".txt", ".md", ".csv", ".pdf", ".xlsx"}


def _parse_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def _parse_csv(path: Path) -> str:
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
        return "\n".join("\t".join(row[:100]) for _, row in zip(range(10000), csv.reader(handle)))


def _parse_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ValueError("PDF parser unavailable") from exc
    return "\n".join((page.extract_text() or "") for page in PdfReader(path).pages[:100])


def _parse_xlsx(path: Path) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("XLSX parser unavailable") from exc
    # 存储层文件名为无扩展名 UUID，openpyxl 会按文件名后缀拒收（InvalidFileException）；
    # 传二进制流则按内容（ZIP 头）校验，与落盘命名无关。
    with path.open("rb") as handle:
        workbook = load_workbook(handle, read_only=True, data_only=True)
        try:
            rows = []
            for sheet in workbook.worksheets[:20]:
                rows.extend(
                    "\t".join("" if value is None else str(value) for value in row[:100])
                    for row in list(sheet.iter_rows(values_only=True))[:10000]
                )
            return "\n".join(rows)
        finally:
            workbook.close()


def parse_file(path: Path, suffix: str, max_chars: int) -> str:
    parsers = {".txt": _parse_text, ".md": _parse_text, ".csv": _parse_csv, ".pdf": _parse_pdf, ".xlsx": _parse_xlsx}
    parser = parsers.get(suffix.lower())
    if parser is None:
        raise ValueError("unsupported file type")
    return parser(path)[:max_chars]